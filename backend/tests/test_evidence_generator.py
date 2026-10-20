import io
import sys
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from fastapi import HTTPException
from fastapi import UploadFile
from bson import ObjectId

from services.evidence_generator import (
    IMAGES, ARCHIVED_CASES, SESSIONS, _reconcile_case_keys, _validated_image,
    build_evidence_docx, download_evidence_report, parse_spreadsheet, update_spreadsheet,
)


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class EvidenceGeneratorTests(unittest.TestCase):
    def test_parser_maps_aliases_and_keeps_every_source_column(self):
        contents = workbook_bytes([
            ["Identificador", "Escenario", "Resultado de ejecución", "Esperado", "Resultado real", "Detalle del error", "Dato adicional"],
            ["CP-51", "Consulta comercial", "Error", "Debe responder 200", "Respondió 500", "Timeout reportado", "Valor original"],
        ])
        cases = parse_spreadsheet(contents, "casos.xlsx")
        self.assertEqual(len(cases), 1)
        case = cases[0]
        self.assertEqual(case["case_id"], "CP-51")
        self.assertEqual(case["name"], "Consulta comercial")
        self.assertEqual(case["status"], "Error")
        self.assertEqual(case["observations"], "Timeout reportado")
        self.assertEqual(case["source_fields"]["Dato adicional"], "Valor original")
        self.assertEqual(case["images"], [])

    def test_parser_does_not_invent_missing_execution_values(self):
        contents = workbook_bytes([
            ["ID caso", "Nombre del caso", "Resultado esperado", "Observaciones"],
            ["CP-01", "Validar alta", "Alta confirmada", None],
        ])
        case = parse_spreadsheet(contents, "casos.xlsx")[0]
        self.assertEqual(case["status"], "")
        self.assertEqual(case["actual_result"], "")
        self.assertEqual(case["observations"], "")

    def test_parser_reads_app_template_without_importing_bug_sheet(self):
        template = Path(__file__).resolve().parents[1] / "data" / "templates" / "casos_template.xlsx"
        workbook = load_workbook(template)
        cases_sheet = workbook["Planilla CP"]
        cases_sheet["B7"] = 51
        cases_sheet["C7"] = "Validar emisión"
        cases_sheet["E7"] = "Póliza emitida"
        cases_sheet["F7"] = "Error"
        cases_sheet["G7"] = "Respondió con rechazo"
        bugs_sheet = workbook["Bugs"]
        bugs_sheet["A2"] = "BUG-1"
        bugs_sheet["B2"] = "Error al emitir"
        bugs_sheet["E2"] = "Abierto"
        output = io.BytesIO()
        workbook.save(output)
        parsed = parse_spreadsheet(output.getvalue(), "casos.xlsx")
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["case_id"], "51")
        self.assertEqual(parsed[0]["name"], "Validar emisión")
        self.assertEqual(parsed[0]["status"], "Error")
        self.assertEqual(parsed[0]["observations"], "Respondió con rechazo")

    def test_parser_keeps_zero_padding_from_excel_cell_format(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ID caso", "Escenario"])
        sheet.append([51, "Consulta"])
        sheet["A2"].number_format = "0000"
        output = io.BytesIO()
        workbook.save(output)
        case = parse_spreadsheet(output.getvalue(), "casos.xlsx")[0]
        self.assertEqual(case["case_id"], "0051")

    def test_reimport_keeps_evidence_keys_when_duplicate_ids_change_order(self):
        original = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso", "Estado"],
            ["CP-01", "Alta", "OK"],
            ["CP-01", "Baja", "Error"],
        ]), "casos.xlsx")
        updated = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso", "Estado"],
            ["CP-01", "Baja", "Error"],
            ["CP-01", "Alta", "OK"],
        ]), "casos.xlsx")
        reconciled, restored = _reconcile_case_keys(updated, original, [])
        self.assertEqual(restored, [])
        self.assertEqual({case["name"]: case["case_key"] for case in reconciled},
                         {case["name"]: case["case_key"] for case in original})

    def test_reimport_restores_archived_case_and_preserves_row_without_id(self):
        original = parse_spreadsheet(workbook_bytes([
            ["Nombre del caso", "Estado"], ["Consulta", "OK"],
        ]), "casos.xlsx")
        moved = parse_spreadsheet(workbook_bytes([
            ["Nombre del caso", "Estado"], ["Otro", "OK"], ["Consulta", "Error"],
        ]), "casos.xlsx")
        reconciled, restored = _reconcile_case_keys(moved, [], original)
        self.assertEqual(restored, [original[0]["case_key"]])
        self.assertEqual(reconciled[1]["case_key"], original[0]["case_key"])
        self.assertNotEqual(reconciled[0]["case_key"], original[0]["case_key"])

    def test_reimport_new_duplicate_does_not_reuse_existing_case_key(self):
        original = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Alta"],
        ]), "casos.xlsx")
        updated = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Nuevo"], ["CP-01", "Alta"],
        ]), "casos.xlsx")
        reconciled, _ = _reconcile_case_keys(updated, original, [])
        self.assertEqual(reconciled[1]["case_key"], original[0]["case_key"])
        self.assertNotEqual(reconciled[0]["case_key"], original[0]["case_key"])

    def test_reimport_preserves_one_renamed_duplicate(self):
        original = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Alta"], ["CP-01", "Baja"],
        ]), "casos.xlsx")
        updated = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Baja"], ["CP-01", "Alta corregida"],
        ]), "casos.xlsx")
        reconciled, _ = _reconcile_case_keys(updated, original, [])
        self.assertEqual(reconciled[0]["case_key"], original[1]["case_key"])
        self.assertEqual(reconciled[1]["case_key"], original[0]["case_key"])

    def test_reimport_recovers_keys_from_old_template_import(self):
        current = parse_spreadsheet(workbook_bytes([
            ["Id. Caso de Prueba", "Objetivo de la prueba", "Validación del resultado"],
            [51, "Validar emisión", "Error"],
        ]), "casos.xlsx")
        legacy = dict(current[0], case_key="legacy-image-key", case_id="Fila 2", name="Fila 2", status="")
        reconciled, restored = _reconcile_case_keys(current, [legacy], [])
        self.assertEqual(reconciled[0]["case_key"], "legacy-image-key")
        self.assertEqual(restored, [])

    def test_oversized_sheet_is_reported_as_413(self):
        with patch("services.evidence_generator.MAX_SPREADSHEET_ROWS", 2):
            with self.assertRaises(HTTPException) as error:
                parse_spreadsheet(workbook_bytes([
                    ["ID caso", "Nombre"], ["CP-01", "Alta"], ["CP-02", "Baja"],
                ]), "casos.xlsx")
        self.assertEqual(error.exception.status_code, 413)

    def test_truncated_image_is_rejected(self):
        image_buffer = io.BytesIO()
        Image.new("RGB", (32, 32)).save(image_buffer, format="PNG")
        with self.assertRaises(HTTPException) as error:
            _validated_image(image_buffer.getvalue()[:-10], "captura.png")
        self.assertEqual(error.exception.status_code, 400)

    def test_report_contains_summary_case_fields_and_embedded_image(self):
        image_buffer = io.BytesIO()
        Image.new("RGB", (320, 180), color=(35, 99, 235)).save(image_buffer, format="PNG")
        image_bytes = image_buffer.getvalue()
        session = {
            "project_name": "Proyecto de prueba",
            "requirement": "REQ-TEST",
            "environment": "QA",
            "source_filename": "casos.xlsx",
            "cases": [{
                "case_key": "case-1",
                "case_id": "CP-01",
                "name": "Alta comercial",
                "status": "Error",
                "source_fields": {
                    "ID": "CP-01",
                    "Caso": "Alta comercial",
                    "Estado": "Error",
                    "Detalle del error": "Respuesta no esperada",
                },
                "images": [{"id": "image-1", "title": "Respuesta del sistema", "caption": "Se observa el error", "category": "Error"}],
            }],
        }
        result = build_evidence_docx(session, {"image-1": image_bytes})
        document = Document(io.BytesIO(result))
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        self.assertIn("Evidencias de ejecución QA", text)
        self.assertIn("CP-01", text)
        self.assertIn("Respuesta no esperada", text)
        self.assertIn("Pend./otros", text)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(len(document.sections), 1)

    def test_report_rejects_missing_stored_image(self):
        session = {
            "project_name": "Proyecto",
            "cases": [{"case_key": "c1", "case_id": "CP-01", "name": "Alta", "images": [{"id": "lost"}]}],
        }
        with self.assertRaisesRegex(ValueError, "Falta la imagen"):
            build_evidence_docx(session, {})

    def test_report_marks_missing_values_and_missing_image(self):
        session = {
            "project_name": "Proyecto",
            "cases": [{"case_key": "case-1", "case_id": "CP-02", "name": "Caso sin datos", "status": "", "source_fields": {"Observaciones": ""}, "images": []}],
        }
        result = build_evidence_docx(session, {})
        document = Document(io.BytesIO(result))
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        self.assertIn("Sin evidencia gráfica adjunta.", text)
        self.assertIn("No informado", text)


class EmptyCursor:
    def __aiter__(self):
        return self

    async def __anext__(self):
        raise StopAsyncIteration


class EvidenceRouteTests(unittest.IsolatedAsyncioTestCase):
    async def test_report_returns_conflict_when_image_is_missing(self):
        session_id = ObjectId()
        image_id = ObjectId()
        raw = {"_id": session_id, "project_name": "QA", "cases": []}
        public = {"project_name": "QA", "cases": [{"images": [{"id": str(image_id)}]}]}
        images_collection = SimpleNamespace(find=MagicMock(return_value=EmptyCursor()))
        with patch("services.evidence_generator._session", new=AsyncMock(return_value=raw)), \
             patch("services.evidence_generator._full_session", new=AsyncMock(return_value=public)), \
             patch("services.evidence_generator._db", return_value={IMAGES: images_collection}):
            with self.assertRaises(HTTPException) as error:
                await download_evidence_report(str(session_id), {"username": "shared", "shared": True})
        self.assertEqual(error.exception.status_code, 409)

    async def test_update_spreadsheet_preserves_images_when_duplicate_ids_move(self):
        original = parse_spreadsheet(workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Alta"], ["CP-01", "Baja"],
        ]), "casos.xlsx")
        updated_bytes = workbook_bytes([
            ["ID caso", "Nombre del caso"], ["CP-01", "Baja"], ["CP-01", "Alta"],
        ])
        session_id = ObjectId()
        raw = {"_id": session_id, "project_name": "QA", "cases": original, "image_count": 2, "archived_case_count": 0}
        archived = SimpleNamespace(
            find=MagicMock(return_value=SimpleNamespace(to_list=AsyncMock(return_value=[]))),
            update_one=AsyncMock(), delete_many=AsyncMock(),
        )
        images_collection = SimpleNamespace(count_documents=AsyncMock(return_value=1))

        async def apply_update(_filter, update):
            raw.update(update["$set"])

        sessions_collection = SimpleNamespace(update_one=AsyncMock(side_effect=apply_update))
        database = {SESSIONS: sessions_collection, ARCHIVED_CASES: archived, IMAGES: images_collection}
        with patch("services.evidence_generator._indexes", new=AsyncMock()), \
             patch("services.evidence_generator._session", new=AsyncMock(side_effect=lambda _: raw)), \
             patch("services.evidence_generator._full_session", new=AsyncMock(side_effect=lambda doc: doc)), \
             patch("services.evidence_generator._db", return_value=database), \
             patch("services.evidence_generator.record_activity", new=AsyncMock()):
            result = await update_spreadsheet(str(session_id), UploadFile(file=io.BytesIO(updated_bytes), filename="nuevo.xlsx"),
                                              {"username": "shared", "shared": True})
        self.assertEqual(result["changes"], {"added": 0, "removed": 0, "restored": 0})
        self.assertEqual({case["name"]: case["case_key"] for case in raw["cases"]},
                         {case["name"]: case["case_key"] for case in original})
        self.assertEqual(raw["image_count"], 2)
        archived.update_one.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()
