import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bson import ObjectId
from docx import Document
from services.auth import current_user
from services.parsing import docx_to_text
from services.quality_records import QualityRecordIn, _payload, _public, update_record


class SharedWorkspaceTests(unittest.IsolatedAsyncioTestCase):
    async def test_access_needs_no_login_or_database(self):
        with patch('services.auth._db', side_effect=AssertionError('No account lookup')):
            user = await current_user()
            stale = await current_user('Bearer expired')
        self.assertTrue(user['shared'])
        self.assertEqual(user, stale)

    async def test_edit_historical_record_keeps_id_and_creator(self):
        oid = ObjectId()
        original = {'_id': oid, 'created_by': 'augusto.giorgi', 'qa_responsible': 'augusto.giorgi', 'generated_cases': 10, 'ok_cases': 8}
        collection = MagicMock()
        collection.find_one = AsyncMock(return_value=dict(original))
        collection.update_one = AsyncMock()
        db = MagicMock()
        db.__getitem__.return_value = collection
        payload = QualityRecordIn(id_req='REQ-1', requirement_name='Consulta', qa_responsible='Augusto Giorgi', design_time_seconds=95, generated_cases=10, ok_cases=8, additional_qa_cases=2)
        with patch('services.quality_records._db', return_value=db):
            response = await update_record(str(oid), payload, await current_user())
        update = collection.update_one.call_args.args[1]['$set']
        self.assertNotIn('created_by', update)
        self.assertNotIn('_id', update)
        self.assertEqual(response['record']['created_by'], 'augusto.giorgi')
        self.assertEqual(response['record']['qa_responsible'], 'Augusto Giorgi')
        self.assertEqual(response['record']['design_time_seconds'], 95)

    def test_legacy_display_is_non_destructive(self):
        source = {'_id': ObjectId(), 'qa_responsible': 'ana.madriz', 'design_time': 1.5}
        result = _public(source)
        self.assertEqual(result['qa_responsible_display'], 'Ana Madriz')
        self.assertEqual(result['design_time_seconds'], 90)
        self.assertEqual(source['qa_responsible'], 'ana.madriz')
        self.assertIn('_id', source)

    def test_responsible_must_be_written(self):
        from fastapi import HTTPException
        payload = QualityRecordIn(id_req='R', requirement_name='R', qa_responsible=' ', design_time_seconds=0, generated_cases=1, ok_cases=1, additional_qa_cases=0)
        with self.assertRaises(HTTPException):
            _payload(payload, {'username': 'shared'})

    def test_docx_preserves_rule_table_order_and_empty_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'rules.docx'
            doc = Document()
            doc.add_paragraph('Regla A')
            table = doc.add_table(rows=1, cols=3)
            table.cell(0, 0).text = 'Campo'
            table.cell(0, 2).text = 'Resultado'
            doc.add_paragraph('Regla B')
            doc.save(path)
            self.assertEqual(docx_to_text(str(path)), 'Regla A\nCampo |  | Resultado\nRegla B')


if __name__ == '__main__':
    unittest.main()
