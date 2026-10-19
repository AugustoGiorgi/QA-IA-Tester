import json
import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from bson import ObjectId

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from services.postman_generator import (  # noqa: E402
    _read_upload,
    build_collection,
    build_environment,
    build_intermediate_model,
    build_readiness,
    analyze_postman_sources,
    download_generated,
    validate_model,
)
from fastapi import HTTPException  # noqa: E402


class AsyncUpload:
    def __init__(self, filename, data):
        self.filename = filename
        self._data = data

    async def read(self, size=-1):
        return self._data if size < 0 else self._data[:size]


class PostmanGeneratorTests(unittest.TestCase):
    def test_individual_requests_excel_cases_and_environment_are_runnable(self):
        import asyncio
        from openpyxl import Workbook

        request = {"name": "Autenticación", "request": {
            "method": "POST", "url": "https://api.example.test/auth", "header": [{"key": "Content-Type", "value": "application/json"}],
            "body": {"mode": "raw", "raw": '{"username":"demo","password":"topsecret","amount":120}'},
        }, "event": [{"listen": "test", "script": {"exec": ["pm.environment.set('sessionId', pm.response.json().id);"]}}]}
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ID", "Nombre del caso", "Endpoint", "Datos de prueba", "Resultado esperado"])
        sheet.append(["CP-01", "Login correcto", "POST /auth", "username=demo; password=topsecret", "HTTP 200"])
        sheet.append(["CP-02", "Login sin contraseña", "POST /auth", "Omitir password", "HTTP 400"])
        buffer = BytesIO()
        workbook.save(buffer)
        sources = [
            {"name": "login.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""},
            asyncio.run(_read_upload(AsyncUpload("casos.xlsx", buffer.getvalue()))),
        ]
        model = build_intermediate_model(sources)
        model["project_name"] = "Login QA"
        collection = build_collection(model)
        environment = {item["key"]: item for item in build_environment(model)["values"]}
        exported = collection["item"][0]["item"]

        self.assertEqual(len(model["endpoints"]), 1)
        self.assertEqual(len(exported), 2)
        self.assertEqual(exported[0]["name"], "CP-01 - Login correcto")
        self.assertIn("{{cp_01_password}}", exported[0]["request"]["body"]["raw"])
        self.assertNotIn("password", exported[1]["request"]["body"]["raw"])
        self.assertEqual(environment["cp_01_password"]["value"], "")
        self.assertEqual(environment["cp_01_amount"]["value"], "120")
        self.assertIn("sessionId", " ".join(exported[0]["event"][0]["script"]["exec"]))
        self.assertIn("pm.response.to.have.status(400)", " ".join(exported[1]["event"][-1]["script"]["exec"]))
        self.assertNotIn("topsecret", json.dumps(collection) + json.dumps(environment))

    def test_missing_queryable_data_is_clear_without_requesting_credentials(self):
        request = {"request": {"method": "GET", "url": "{{baseUrl}}/polizas/:policyNumber?password={{password}}"}}
        model = build_intermediate_model([{"name": "poliza.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""}])
        model["project_name"] = "Polizas"
        readiness = build_readiness(model)
        missing = {item["key"] for item in readiness["missing_data"]}
        self.assertIn("policyNumber", missing)
        self.assertIn("password", missing)
        self.assertIn("policyNumber", readiness["query_request"])
        self.assertNotIn("password", readiness["query_request"])

    def test_multiple_base_urls_and_functional_document_do_not_add_fake_endpoints(self):
        first = {"request": {"method": "GET", "url": "https://api.one.test/polizas"}}
        second = {"request": {"method": "GET", "url": "https://api.two.test/clientes"}}
        sources = [
            {"name": "polizas.json", "extension": ".json", "text": json.dumps(first), "size": 1, "warning": ""},
            {"name": "clientes.json", "extension": ".json", "text": json.dumps(second), "size": 1, "warning": ""},
            {"name": "funcional.docx", "extension": ".docx", "text": "Regla: GET /polizas requiere póliza vigente. GET /otra-ruta no es un endpoint.", "size": 1, "warning": ""},
        ]
        model = build_intermediate_model(sources)
        self.assertEqual(len(model["endpoints"]), 2)
        self.assertEqual(len({item["id"] for item in model["endpoints"]}), 2)
        urls = [item["request"]["url"]["raw"] for folder in build_collection(model)["item"] for item in folder["item"]]
        self.assertTrue(any("{{baseUrl}}" in url for url in urls))
        self.assertTrue(any("{{baseUrl2}}" in url for url in urls))
        self.assertTrue(any("requiere póliza vigente" in item["request"]["description"] for folder in build_collection(model)["item"] for item in folder["item"]))

    def test_identical_paths_on_different_hosts_remain_separate(self):
        sources = [
            {"name": name, "extension": ".json", "text": json.dumps({"request": {"method": "GET", "url": f"https://{host}/health"}}), "size": 1, "warning": ""}
            for name, host in (("a.json", "api.one.test"), ("b.json", "api.two.test"))
        ]
        model = build_intermediate_model(sources)
        self.assertEqual(len(model["endpoints"]), 2)
        self.assertEqual(len({item["id"] for item in model["endpoints"]}), 2)
        cases = {"name": "casos.csv", "extension": ".csv", "text": "CASO CP-01 - Consultar\nEndpoint: GET https://api.two.test/health\nResultado esperado: HTTP 200\n", "size": 1, "warning": "", "is_cases_file": True}
        model = build_intermediate_model(sources + [cases])
        target = next(item for item in model["endpoints"] if "api.two.test" in item.get("original_base_url", ""))
        self.assertEqual(model["associations"][0]["endpoint_id"], target["id"])

    def test_two_collections_with_same_base_variable_use_distinct_hosts(self):
        sources = []
        for index, host in enumerate(("api.one.test", "api.two.test"), start=1):
            collection = {
                "info": {"name": f"API {index}", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
                "variable": [{"key": "baseUrl", "value": f"https://{host}"}],
                "item": [{"name": "Consultar", "request": {"method": "GET", "url": "{{baseUrl}}/health"}}],
            }
            sources.append({"name": f"api{index}.json", "extension": ".json", "text": json.dumps(collection), "size": 1, "warning": ""})
        model = build_intermediate_model(sources)
        self.assertEqual(len(model["endpoints"]), 2)
        urls = [item["request"]["url"]["raw"] for folder in build_collection(model)["item"] for item in folder["item"]]
        self.assertTrue(any("{{baseUrl}}" in url for url in urls))
        self.assertTrue(any("{{baseUrl2}}" in url for url in urls))
        environment = {item["key"]: item["value"] for item in build_environment(model)["values"]}
        self.assertEqual(environment["baseUrl"], "https://api.one.test")
        self.assertEqual(environment["baseUrl2"], "https://api.two.test")

    def test_case_route_prefers_literal_over_parameterized_route(self):
        collection = {
            "info": {"name": "API", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
            "item": [
                {"name": "Total", "request": {"method": "GET", "url": "https://api.test/polizas/total"}},
                {"name": "Detalle", "request": {"method": "GET", "url": "https://api.test/polizas/:id"}},
            ],
        }
        cases = ("CASO CP-01 - Consultar total\nEndpoint: GET /polizas/total\nResultado esperado: HTTP 200\n\n"
                 "CASO CP-02 - Consultar detalle\nEndpoint: GET /polizas/:id\nResultado esperado: HTTP 200\n")
        model = build_intermediate_model([
            {"name": "api.json", "extension": ".json", "text": json.dumps(collection), "size": 1, "warning": ""},
            {"name": "casos.csv", "extension": ".csv", "text": cases, "size": 1, "warning": "", "is_cases_file": True},
        ])
        endpoints = {item["id"]: item["path"] for item in model["endpoints"]}
        self.assertEqual(endpoints[model["associations"][0]["endpoint_id"]], "/polizas/total")
        self.assertEqual(endpoints[model["associations"][1]["endpoint_id"]], "/polizas/{id}")

    def test_csv_case_file_matches_one_request(self):
        import asyncio
        csv_bytes = b"ID;Nombre del caso;Endpoint;Datos de prueba;Resultado esperado\nCP-09;Consultar cliente;GET /clientes;clientId=42;HTTP 200\n"
        source = asyncio.run(_read_upload(AsyncUpload("casos.csv", csv_bytes)))
        request = {"request": {"method": "GET", "url": "https://api.test/clientes?clientId={{clientId}}"}}
        model = build_intermediate_model([{"name": "cliente.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""}, source])
        self.assertEqual(len(model["test_cases"]), 1)
        self.assertEqual(len(build_collection(model)["item"][0]["item"]), 1)

    def test_bom_json_and_multiline_excel_input_keep_case_values(self):
        import asyncio
        from openpyxl import Workbook

        request = {"request": {"method": "POST", "url": "https://api.test/polizas", "body": {
            "mode": "raw", "raw": '{"policyNumber":"example","amount":0}',
        }}}
        workbook = Workbook()
        workbook.active.append(["ID", "Nombre del caso", "Endpoint", "Datos de prueba", "Resultado esperado"])
        workbook.active.append(["CP-01", "Crear póliza", "POST /polizas",
                                '{\n  "policyNumber": "123",\n  "amount": 10\n}', "HTTP 201"])
        buffer = BytesIO()
        workbook.save(buffer)
        sources = [
            asyncio.run(_read_upload(AsyncUpload("request.json", b"\xef\xbb\xbf" + json.dumps(request).encode()))),
            asyncio.run(_read_upload(AsyncUpload("casos.xlsx", buffer.getvalue()))),
        ]
        model = build_intermediate_model(sources)
        self.assertEqual(len(model["endpoints"]), 1)
        self.assertEqual(len(model["test_cases"]), 1)
        self.assertEqual(len(build_collection(model)["item"][0]["item"]), 1)
        environment = {item["key"]: item["value"] for item in build_environment(model)["values"]}
        self.assertEqual(environment["cp_01_amount"], "10")
        self.assertEqual(environment["cp_01_policyNumber"], "123")

    def test_case_overrides_boolean_and_null_body_values(self):
        request = {"request": {"method": "POST", "url": "https://api.test/flags", "body": {
            "mode": "raw", "raw": '{"active":true,"reference":"sample"}',
        }}}
        cases = (
            'CASO CP-01 - Desactivar\nEndpoint: POST /flags\nDatos de prueba: {"active":false}\nResultado esperado: HTTP 200\n\n'
            'CASO CP-02 - Enviar nulo\nEndpoint: POST /flags\nDatos de prueba: {"active":null}\nResultado esperado: HTTP 400\n'
        )
        model = build_intermediate_model([
            {"name": "flags.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""},
            {"name": "casos.csv", "extension": ".csv", "text": cases, "size": 1, "warning": "", "is_cases_file": True},
        ])
        items = build_collection(model)["item"][0]["item"]
        self.assertEqual(len(items), 2)
        self.assertIs(json.loads(items[0]["request"]["body"]["raw"])["active"], False)
        self.assertIsNone(json.loads(items[1]["request"]["body"]["raw"])["active"])

    def test_bearer_auth_is_parameterized_without_exporting_literal(self):
        request = {"request": {"method": "GET", "url": "https://api.test/clientes", "auth": {
            "type": "bearer", "bearer": [{"key": "token", "value": "literal-secret-token"}],
        }}}
        model = build_intermediate_model([{"name": "request.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""}])
        collection = build_collection(model)
        environment = build_environment(model)
        self.assertNotIn("literal-secret-token", json.dumps(collection) + json.dumps(environment))
        self.assertIn("bearerToken", {item["key"] for item in environment["values"]})

    def test_collection_auth_and_prerequest_script_survive_export(self):
        collection = {
            "info": {"name": "API", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
            "auth": {"type": "bearer", "bearer": [{"key": "token", "value": "raw-token-secret"}]},
            "event": [{"listen": "prerequest", "script": {"type": "text/javascript", "exec": ["pm.variables.set('timestamp', Date.now());"]}}],
            "item": [{"name": "Consultar", "request": {"method": "GET", "url": "https://api.test/clientes"}}],
        }
        model = build_intermediate_model([{"name": "api.json", "extension": ".json", "text": json.dumps(collection), "size": 1, "warning": ""}])
        item = build_collection(model)["item"][0]["item"][0]
        self.assertEqual(item["request"]["auth"]["bearer"][0]["value"], "{{bearerToken}}")
        self.assertEqual(item["event"][0]["listen"], "prerequest")


class PostmanAnalyzeTests(unittest.IsolatedAsyncioTestCase):
    async def test_analyze_saves_exact_case_and_pending_data(self):
        source = json.dumps({"request": {"method": "GET", "url": "https://api.test/polizas/:policyNumber"}}).encode()
        cases = b"ID;Nombre del caso;Endpoint;Resultado esperado\nCP-01;Consultar poliza;GET /polizas/:policyNumber;HTTP 200\n"
        fake_collection = SimpleNamespace(insert_one=AsyncMock(return_value=SimpleNamespace(inserted_id=ObjectId())))
        with patch("services.postman_generator._db", return_value={"PostmanGenerationDrafts": fake_collection}), \
                patch("services.postman_generator.record_activity", new=AsyncMock()):
            response = await analyze_postman_sources(
                files=[AsyncUpload("endpoint.json", source)], test_cases_file=AsyncUpload("casos.csv", cases),
                manual_text="", project_name="Pólizas QA", user={"username": "shared", "role": "qa", "shared": True},
            )
        model = response["draft"]["model"]
        self.assertEqual(model["readiness"]["request_count"], 1)
        self.assertIn("policyNumber", {item["key"] for item in model["readiness"]["missing_data"]})
        self.assertIn("No ejecutes la consulta", model["readiness"]["query_request"])
        fake_collection.insert_one.assert_awaited_once()

    async def test_unmatched_excel_case_is_rejected_before_saving(self):
        source = json.dumps({"request": {"method": "GET", "url": "https://api.test/clientes"}}).encode()
        cases = b"ID;Nombre del caso;Endpoint;Resultado esperado\nCP-01;Emitir poliza;POST /polizas;HTTP 201\n"
        with self.assertRaises(HTTPException) as raised:
            await analyze_postman_sources(
                files=[AsyncUpload("endpoint.json", source)], test_cases_file=AsyncUpload("casos.csv", cases),
                manual_text="", project_name="QA", user={"username": "shared", "role": "qa", "shared": True},
            )
        self.assertEqual(raised.exception.status_code, 422)
        self.assertIn("CP-01", raised.exception.detail)

    async def test_bundle_contains_only_collection_and_environment(self):
        request = {"request": {"method": "GET", "url": "https://api.test/clientes"}}
        model = build_intermediate_model([{"name": "request.json", "extension": ".json", "text": json.dumps(request), "size": 1, "warning": ""}])
        model["project_name"] = "QA"
        doc = {"project_name": "QA", "model": model}
        with patch("services.postman_generator._get_draft", new=AsyncMock(return_value=doc)), patch("services.postman_generator.record_activity", new=AsyncMock()):
            response = await download_generated("bundle", "draft-id", {"username": "shared", "role": "qa", "shared": True})
        payload = b"".join([chunk async for chunk in response.body_iterator])
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(set(archive.namelist()), {"collection.json", "environment.json"})
            self.assertTrue(json.loads(archive.read("collection.json"))["item"])

    def test_rejects_oversized_source_before_parsing(self):
        import asyncio

        with patch("services.postman_generator.MAX_SOURCE_BYTES", 4), self.assertRaises(HTTPException) as raised:
            asyncio.run(_read_upload(AsyncUpload("collection.json", b"12345")))
        self.assertEqual(raised.exception.status_code, 413)

    def test_limits_expanded_spreadsheet_rows(self):
        import asyncio
        from openpyxl import Workbook

        workbook = Workbook()
        for value in ("ID", "CP-01", "CP-02"):
            workbook.active.append([value])
        buffer = BytesIO()
        workbook.save(buffer)
        with patch("services.postman_generator.MAX_SPREADSHEET_ROWS", 2), self.assertRaises(HTTPException) as raised:
            asyncio.run(_read_upload(AsyncUpload("casos.xlsx", buffer.getvalue())))
        self.assertEqual(raised.exception.status_code, 413)

    def test_openapi_users_domain_generates_collection(self):
        openapi = """
openapi: 3.0.0
servers:
  - url: https://api.example.test
paths:
  /users/{userId}:
    get:
      summary: Obtener usuario
      parameters:
        - name: userId
          in: path
          required: true
      responses:
        "200":
          description: Usuario encontrado
"""
        model = build_intermediate_model(
            [{"name": "users.yaml", "extension": ".yaml", "text": openapi, "size": len(openapi), "warning": ""}],
            "CASO 1: Obtener usuario existente\nResultado esperado: HTTP 200",
        )
        self.assertEqual(len(model["endpoints"]), 1)
        self.assertEqual(model["endpoints"][0]["method"], "GET")
        collection = build_collection(model)
        self.assertEqual(collection["info"]["schema"], "https://schema.getpostman.com/json/collection/v2.1.0/collection.json")
        self.assertEqual(collection["item"][0]["item"][0]["request"]["method"], "GET")

    def test_curl_inventory_domain_detects_secret_and_environment(self):
        text = "curl -X POST https://stock.example.test/items -H 'X-API-Key: abc123456789' -H 'Content-Type: application/json' -d '{\"sku\":\"A1\"}'"
        model = build_intermediate_model(
            [{"name": "inventory.txt", "extension": ".txt", "text": text, "size": len(text), "warning": ""}]
        )
        self.assertEqual(model["endpoints"][0]["method"], "POST")
        self.assertTrue(any(w["type"] == "possible_secret" for w in model["warnings"]))
        environment = build_environment(model)
        keys = {item["key"] for item in environment["values"]}
        self.assertIn("baseUrl", keys)

    def test_file_processing_domain_keeps_unmatched_case_visible(self):
        docs = """
CASO FP-01 Procesar archivo invalido
Pasos:
1. Enviar archivo corrupto al proceso batch.
Resultado esperado: Rechazo controlado.
"""
        model = build_intermediate_model(
            [{"name": "file-cases.md", "extension": ".md", "text": docs, "size": len(docs), "warning": ""}]
        )
        self.assertEqual(len(model["test_cases"]), 1)
        self.assertEqual(model["associations"][0]["confidence"], "Sin coincidencia")
        validation = validate_model(model)
        self.assertTrue(validation["valid"])
        self.assertIn("collection", json.dumps(build_collection(model)))

    def test_xlsx_cases_file_is_readable(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Casos"
        sheet.append(["ID", "Nombre", "Resultado esperado"])
        sheet.append(["CP-01", "Crear siniestro", "HTTP 201"])
        sheet.append(["CP-02", "Validar cero", 0])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = AsyncUpload("casos.xlsx", buffer.getvalue())

        import asyncio

        source = asyncio.run(_read_upload(upload))
        self.assertEqual(source["extension"], ".xlsx")
        self.assertIn("Crear siniestro", source["text"])
        self.assertIn("Resultado esperado: 0", source["text"])

    def test_cases_spreadsheet_exports_only_its_named_cases(self):
        from openpyxl import Workbook

        collection = {
            "info": {"name": "API", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
            "item": [
                {"name": "Autenticacion", "request": {"method": "POST", "url": {"raw": "https://api.example.test/api/auth"}}},
                {"name": "Consultar", "request": {"method": "GET", "url": {"raw": "https://api.example.test/api/customer"}}},
            ],
        }
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Matriz de casos de prueba"])
        sheet.append(["ID", "Nombre del caso", "Endpoint", "Categoria", "Precondicion", "Datos de prueba", "Pasos", "Resultado esperado"])
        sheet.append(["CP-24", "CP-24 - Autenticacion sin contraseña", "POST /api/auth", "Negativo", "Usuario habilitado", "Omitir password", "Enviar request", "Rechazo controlado"])
        detail = workbook.create_sheet("Detalle Requests")
        detail.append(["Detalle exacto de requests por caso"])
        detail.append(["ID", "Request", "Metodo", "Endpoint"])
        detail.append(["CP-24", "1. Autenticar sin contraseña", "POST", "/api/auth"])
        buffer = BytesIO()
        workbook.save(buffer)

        import asyncio

        collection_source = {"name": "collection.json", "extension": ".json", "text": json.dumps(collection), "size": 1, "warning": ""}
        cases_source = asyncio.run(_read_upload(AsyncUpload("casos.xlsx", buffer.getvalue())))
        model = build_intermediate_model([collection_source, cases_source])
        requests = build_collection(model)["item"][0]["item"]

        self.assertEqual(len(model["test_cases"]), 1)
        self.assertEqual(len(requests), 1)
        self.assertEqual(requests[0]["name"], "CP-24 - Autenticacion sin contraseña")
        self.assertEqual(requests[0]["request"]["method"], "POST")

    def test_endpoint_without_excel_gets_base_ok_and_base_error_cases(self):
        openapi = """
openapi: 3.0.0
paths:
  /polizas/{id}:
    post:
      summary: Actualizar poliza
      requestBody:
        content:
          application/json:
            example:
              estado: activa
      responses:
        "200":
          description: OK
"""
        model = build_intermediate_model(
            [{"name": "polizas.yaml", "extension": ".yaml", "text": openapi, "size": len(openapi), "warning": ""}]
        )
        generated = [case for case in model["test_cases"] if case.get("generated")]
        self.assertEqual(len(generated), 2)
        self.assertTrue(any("caso base OK" in case["name"] for case in generated))
        self.assertTrue(any("caso base con error" in case["name"] for case in generated))

    def test_business_rule_case_is_generated_only_when_documented(self):
        openapi = """
openapi: 3.0.0
paths:
  /cotizaciones:
    post:
      summary: Crear cotizacion
      description: Debe validar regla de negocio de monto limite y estado vigente.
      responses:
        "200":
          description: OK
"""
        model = build_intermediate_model(
            [{"name": "cotizaciones.yaml", "extension": ".yaml", "text": openapi, "size": len(openapi), "warning": ""}]
        )
        generated = [case for case in model["test_cases"] if case.get("generated")]
        self.assertEqual(len(generated), 3)
        self.assertTrue(any("logica de negocio" in case["name"] for case in generated))

    def test_postman_collection_repeated_requests_count_as_unique_endpoint(self):
        collection = {
            "info": {"schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
            "item": [
                {"name": "CP-01", "item": [
                    {"name": "Consultar asegurado OK", "request": {"method": "GET", "url": {"raw": "{{baseUrl}}/api/v1/Global/consultarasegurado?cuit={{cuit}}"}}},
                ]},
                {"name": "CP-02", "item": [
                    {"name": "Consultar asegurado error", "request": {"method": "GET", "url": {"raw": "{{baseUrl}}/api/v1/Global/consultarasegurado?cuit={{cuitInvalido}}"}}},
                ]},
            ],
        }
        text = json.dumps(collection)
        model = build_intermediate_model(
            [{"name": "collection.json", "extension": ".json", "text": text, "size": len(text), "warning": ""}]
        )
        self.assertEqual(len(model["endpoints"]), 1)

    def test_postman_collection_variables_export_to_environment(self):
        collection = {
            "info": {"schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
            "variable": [
                {"key": "baseUrl", "value": "https://api.example.test"},
                {"key": "accessToken", "value": "do-not-export-this", "type": "secret"},
            ],
            "item": [{"name": "Consultar", "request": {
                "method": "GET",
                "url": {"raw": "{{baseUrl}}/users/:userId?include=profile"},
            }}],
        }
        text = json.dumps(collection)
        model = build_intermediate_model([
            {"name": "api.postman_collection.json", "extension": ".json", "text": text, "size": len(text), "warning": ""}
        ])

        request_url = build_collection(model)["item"][0]["item"][0]["request"]["url"]["raw"]
        environment = {item["key"]: item for item in build_environment(model)["values"]}

        self.assertIn("{{userId}}", request_url)
        self.assertEqual(environment["baseUrl"]["value"], "https://api.example.test")
        self.assertEqual(environment["userId"]["value"], "")
        self.assertEqual(environment["accessToken"]["value"], "")


if __name__ == "__main__":
    unittest.main()
